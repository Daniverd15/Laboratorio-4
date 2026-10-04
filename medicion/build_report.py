#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Genera Informe_Laboratorio4.docx (2 paginas A4) con python-docx.
Uso: python medicion/build_report.py   (requiere python-docx; el PDF se obtiene con Word o LibreOffice)"""
import os
import re
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
EV = os.path.join(RAIZ, "evidencias")
SALIDA = os.path.join(RAIZ, "Informe_Laboratorio4.docx")

AZUL = RGBColor(0x1F, 0x4E, 0x9E)
GRIS = RGBColor(0x55, 0x5B, 0x66)
VERDE = RGBColor(0x1E, 0x7E, 0x4B)
ROJO = RGBColor(0xB0, 0x3A, 0x2E)
FUENTE = "Calibri"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.2)
sec.top_margin, sec.bottom_margin = Cm(1.0), Cm(0.9)
ANCHO = 21.0 - 2.4

st = doc.styles["Normal"]
st.font.name = FUENTE
st.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
st.font.size = Pt(8)
st.paragraph_format.space_after = Pt(1)
st.paragraph_format.space_before = Pt(0)
st.paragraph_format.line_spacing = 1.0


def runs(p, texto, negrita=False, tam=None, color=None, cursiva=False):
    """Formato en linea: **negrita**, `codigo` (Consolas), *cursiva*."""
    for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)", texto):
        if not tok:
            continue
        neg, cur, cod = negrita, cursiva, False
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            tok, neg = tok[2:-2], True
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            tok, cod = tok[1:-1], True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            tok, cur = tok[1:-1], True
        r = p.add_run(tok)
        r.bold, r.italic = neg, cur
        if cod:
            r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        if tam:
            r.font.size = Pt(tam - 0.7 if cod else tam)
        if color is not None:
            r.font.color.rgb = color
    return p


def parrafo(texto="", negrita=False, tam=None, color=None, alinear=None, antes=0, despues=1, contenedor=None, cursiva=False):
    p = (contenedor or doc).add_paragraph()
    p.paragraph_format.space_before = Pt(antes)
    p.paragraph_format.space_after = Pt(despues)
    if alinear:
        p.alignment = alinear
    if texto:
        runs(p, texto, negrita, tam, color, cursiva)
    return p


def titulo(texto):
    p = parrafo(texto, negrita=True, tam=9.5, color=AZUL, antes=3, despues=1)
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "4"), ("space", "1"), ("color", "1F4E9E")):
        bt.set(qn("w:" + k), v)
    b.append(bt)
    pPr.append(b)
    return p


def vinneta(texto, contenedor=None, tam=None):
    p = parrafo("", contenedor=contenedor, despues=0.5)
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.first_line_indent = Cm(-0.3)
    runs(p, "• " + texto, tam=tam)
    return p


def sombra(celda, hexcolor):
    tcPr = celda._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hexcolor)
    tcPr.append(sh)


def margenes(celda, arriba=15, abajo=15, izq=40, der=40):
    tcPr = celda._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for k, v in (("top", arriba), ("left", izq), ("bottom", abajo), ("right", der)):
        e = OxmlElement("w:" + k)
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tcPr.append(m)


def bordes(tabla, color="B8BEC9", sz="4"):
    b = OxmlElement("w:tblBorders")
    for k in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + k)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), sz)
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        b.append(e)
    tabla._tbl.tblPr.append(b)


def sin_bordes(t):
    b = OxmlElement("w:tblBorders")
    for k in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + k)
        e.set(qn("w:val"), "nil")
        b.append(e)
    t._tbl.tblPr.append(b)


def tabla(filas, anchos, tam=7, zebra=True):
    t = doc.add_table(rows=len(filas), cols=len(anchos))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    bordes(t)
    for i, fila in enumerate(filas):
        for j, txt in enumerate(fila):
            c = t.cell(i, j)
            c.width = Cm(anchos[j])
            margenes(c)
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            es_cab = i == 0
            runs(p, txt, negrita=es_cab or j == 0, tam=tam, color=RGBColor(255, 255, 255) if es_cab else None)
            if es_cab:
                sombra(c, "1F4E9E")
            elif zebra and i % 2 == 0:
                sombra(c, "EEF3FB")
    return t


def imagen(celda, ruta, ancho_cm, pie=None, tam_pie=6.5):
    p = celda.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(ruta, width=Cm(ancho_cm))
    if pie:
        q = celda.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.space_after = Pt(0)
        runs(q, pie, tam=tam_pie, color=GRIS, cursiva=True)


def recorte(nombre, caja):
    """Recorta la zona relevante de una captura para que el texto sea legible en el informe.
    El archivo original de evidencias/ no se modifica; el recorte va a una carpeta temporal."""
    import tempfile
    from PIL import Image
    ruta = os.path.join(EV, nombre)
    if not caja:
        return ruta
    d = os.path.join(tempfile.gettempdir(), "lab4_recortes")
    os.makedirs(d, exist_ok=True)
    out = os.path.join(d, os.path.splitext(nombre)[0] + "_recorte.png")
    Image.open(ruta).crop(caja).save(out)
    return out


def cuadricula(items, cols, ancho_total, tam_pie=7):
    filas = (len(items) + cols - 1) // cols
    tt = doc.add_table(rows=filas, cols=cols)
    tt.autofit = False
    sin_bordes(tt)
    w = ancho_total / cols
    for k, it in enumerate(items):
        r, pie = it[0], it[1]
        caja = it[2] if len(it) > 2 else None
        c = tt.cell(k // cols, k % cols)
        c.width = Cm(w)
        margenes(c, 10, 10, 15, 15)
        imagen(c, recorte(r, caja), w - 0.3, pie, tam_pie=tam_pie)
    return tt


# ------------------------------------------------------------------ PAGINA 1
parrafo("Laboratorio 4 — AMQP + tercer protocolo (HTTPS) frente a MQTT, hacia Azure IoT Central", negrita=True, tam=12.5, color=AZUL, despues=0)
parrafo("IoT + Cloud + Sistemas Distribuidos · UNAB · Escenario UNAB-Ambiental · App IoT Central: lab4-medidordeclima-unab · "
        "Repositorio: https://github.com/Daniverd15/Laboratorio-4", tam=7.5, color=GRIS, despues=0)
parrafo("Integrantes: Daniel Villamizar · Tomás Urieles · David Guerrero", negrita=True, tam=8.5, despues=2)

titulo("1. Qué se hizo y rol de AMQP en Azure")
parrafo("Se publicaron las **mismas 3 variables** (temperature, humedad, illuminance; 53 B de JSON) a la **misma app de IoT Central** por tres transportes desde una VM Azure "
        "(Ubuntu, Python 3.10): **MQTT** (línea base del Lab 3), **AMQP 1.0** (Etapa 1) y **HTTPS/REST** (Etapa 2, tercer protocolo). "
        "AMQP 1.0 es el protocolo del **plano de servicio** de Azure (Service Bus, Event Hubs; el endpoint compatible de IoT Hub se lee por AMQP): "
        "conexión → sesión → **enlaces** con **crédito** (control de flujo) y **liquidación** por mensaje (accepted/rejected/released), "
        "frente al pub/sub ligero de MQTT. IoT Hub también lo ofrece al dispositivo (5671, o 443 por WebSockets). "
        "La suscripción de los Labs 1–3 se quedó sin créditos: se **recreó** la app, la plantilla Hobo MX-100 v3, la VM y los dispositivos en una suscripción nueva.", tam=8.5)

t = doc.add_table(rows=1, cols=2)
t.autofit = False
sin_bordes(t)
c1, c2 = t.cell(0, 0), t.cell(0, 1)
c1.width, c2.width = Cm(10.6), Cm(ANCHO - 10.6)
for c in (c1, c2):
    margenes(c, 0, 0, 20, 20)
imagen(c1, os.path.join(EV, "00-arquitectura-flujo.png"), 10.4, "Fig. 1 — Tres caminos hacia IoT Central; DPS por HTTPS en todos.")
c2.text = ""
parrafo("Etapa 1 — AMQP hacia IoT Central (checkpoint cumplido)", negrita=True, tam=8.5, color=AZUL, contenedor=c2, despues=1)
for s_ in [
    "**Cliente:** Apache Qpid **Proton** 0.40. El SDK `azure-iot-device` de **Python no tiene AMQP** (0 archivos lo mencionan; su `websockets` es MQTT-WS).",
    "**Flujo:** TLS **5671** verificado + SASL ANONYMOUS → `open/begin` → nodo `$cbs`: `put-token` con el **SAS firmado a mano** (HMAC-SHA256) → `status-code 200` (10 ms) → enlace `/devices/{id}/messages/events` con **crédito 50** → `transfer` → `disposition accepted` (≈103 ms).",
    "**Anotado (lo que pide el lab):** puerto **5671** con TLS verificado; **sí aparece** en datos sin procesar (Fig. 2), dashboard (Fig. 3) y Explorador de datos (Fig. 5), dispositivo *Conectado*. "
    "**Qué cambia vs MQTT:** conexión persistente en ambos, pero AMQP añade **sesión + 3 enlaces** (`$cbs` ×2 y telemetría); tamaño percibido **438 B vs 365 B** por mensaje (+20 %); ack por mensaje; "
    "**SAS renovable en caliente** (TTL 120 s: 3 renovaciones en 330 s, **0 reconexiones**); debug por tramas (`PN_TRACE_FRM`, Fig. 6).",
    "**Qué resuelve mejor que MQTT:** contrapresión por crédito, liquidación explícita y seguridad en banda. **No** gana en latencia ni en bytes (Sec. 3).",
    "**Tropiezos reales:** el wheel de Proton para Linux **no trae TLS** (se compiló contra OpenSSL); Proton entrega objetos nuevos por evento (los enlaces se comparan por nombre).",
]:
    vinneta(s_, c2, tam=7.5)

titulo("2. Etapa 2 — Tercer protocolo: HTTPS (REST) hacia IoT Hub / Central")
parrafo("**Justificación de negocio:** un sensor de batería que despierta cada pocos minutos paga más por mantener una sesión que por una petición, y HTTPS (443) atraviesa los firewalls/proxys del cliente sin abrir puertos. "
        "**Implementación** (`https_telemetria.py`, solo stdlib, 117 líneas): `POST https://{hub}/devices/{id}/messages/events` con `Authorization: SAS` (firmado a mano), `Content-Type: application/json` y `Content-Encoding: utf-8` → `204`. "
        "Hallazgo: sin ese último encabezado el Hub acepta el POST pero **Central no interpreta el JSON**. Depurable con `curl` (204 válido / 401 `IotHubUnauthorizedAccess` con firma alterada). "
        "También se simuló en el **ESP32 (Wokwi)**: DPS por REST + POST con `HTTPClient` (1.er POST ≈ 4,0 s por el handshake TLS del ESP32; siguientes ≈ 162 ms) y el testigo MQTT (`TCP+TLS+CONNECT` ≈ 3,8 s); ambos llegan a Central (Fig. 8). "
        "**Checkpoint (log + captura):** log del cliente en la VM con `204` en cada envío (Fig. 7) y la telemetría en Central (Fig. 4).", tam=8.5)
t3 = doc.add_table(rows=1, cols=2)
t3.autofit = False
sin_bordes(t3)
a_, b_ = t3.cell(0, 0), t3.cell(0, 1)
a_.width = b_.width = Cm(ANCHO / 2)
for c in (a_, b_):
    margenes(c, 0, 0, 20, 20)
parrafo("Ventajas", negrita=True, tam=8, contenedor=a_, color=VERDE, despues=0)
for s_ in ["Sin librería ni broker; **117 líneas**; depurable con `curl` y códigos de estado estándar.", "443 pasa casi cualquier firewall/proxy.", "Funciona en el ESP32 simulado (HTTPClient)."]:
    vinneta(s_, a_, tam=7.5)
parrafo("Desventajas (medidas; † = documentación)", negrita=True, tam=8, contenedor=b_, color=ROJO, despues=0)
for s_ in ["**901 B/msg** (17×); sin keep-alive **5,7 KB y 199 ms** por mensaje.", "El SAS viaja en **cada** petición; Central lo muestra **Desconectado** (sin presencia, Fig. 4).", "Sin comandos hacia el dispositivo† (habría que sondear mensajes C2D)."]:
    vinneta(s_, b_, tam=7.5)

titulo("3. Mediciones (VM Azure → IoT Hub; N=51 mensajes de 53 B; bytes IP totales con contadores iptables)")
tabla([
    ["Protocolo", "Setup hasta enviar", "ACK p50 / p95", "Bytes/msg en cable", "Conexión + 1 msg", "Paq. /msg", "Corte 16 s (reset): perdidos · 1.er ACK tras volver la red", "Corte 16 s (drop)"],
    ["MQTT 8883 QoS 1", "45 ms", "103 / 118 ms", "365 (6,9×)", "6 180 B", "4,0", "0 · 1,5 s · paho reencola solo", "0 pérd.; **sin reconectar** (TCP retransmite), 0,3 s"],
    ["MQTT 8883 QoS 0", "50 ms", "— (sin ack)", "221 (4,2×)", "6 129 B", "1,9", "(QoS 0 no garantiza entrega)", "—"],
    ["MQTT-WS 443 QoS 1", "166 ms", "105 / 113 ms", "370 (7,0×)", "11 922 B", "4,0", "—", "—"],
    ["**AMQP 5671**", "100 ms", "103 / 111 ms", "**438 (8,3×)**", "9 162 B", "4,2", "0 · 1,6 s · **1 reenvío** (posible duplicado)", "0 pérd.; **sin reconectar**, 0,2 s"],
    ["**HTTPS 443** keep-alive", "41 ms", "109 / 115 ms", "**901 (17×)**", "5 730 B", "5,0", "0 · 1,3 s · cola y reintentos en la app", "0 pérd.; reconecta tras timeout de 20 s, 5,3 s"],
    ["HTTPS 443 sin keep-alive", "35 ms", "199 / 236 ms", "5 699 (107×)", "5 691 B", "16,1", "—", "—"],
], [3.2, 1.5, 1.9, 2.1, 1.7, 1.0, 3.3, 2.9], tam=7)
parrafo("Lectura: la **latencia no distingue** a MQTT, AMQP y HTTPS con keep-alive (el ack tarda ≈100 ms en los tres, mientras un `put-token` AMQP se contesta en 10 ms: domina el Hub); lo que cambia es el **overhead** (MQTT < AMQP < HTTPS), el costo de **abrir** la conexión (AMQP hace SASL+open+begin+CBS+attach ≈ 5 viajes) y **quién** garantiza la entrega "
        "(paho en MQTT QoS 1; la aplicación en AMQP/HTTPS: cola en memoria + backoff ≤ 2 s). Con AMQP hay que **deduplicar por message-id**. "
        "En MQTT el SAS solo existe en el CONNECT: en 330 s con TTL 120 s el Hub **no** cortó la conexión, pero no debe confiarse en eso. "
        "Evidencia de las mediciones: `evidencias/05` (bytes y latencia), `06` (corte de red), `07` (renovación del SAS) y los CSV/logs de `mediciones/`.", tam=7.5, despues=2, antes=2)

titulo("4. Repositorio entregable y límites")
parrafo("**Repo:** https://github.com/Daniverd15/Laboratorio-4 — `amqp/amqp_iothub.py` (script AMQP en la VM), `https/https_telemetria.py` (tercer protocolo), `mqtt/` (línea base), `wokwi/` (ESP32 MQTT y HTTPS), "
        "`medicion/` (benchmarks y corte de red), `README.md` (dependencias y cómo reproducir, **sin secretos**: claves por variables de entorno) y `/evidencias` (logs, tramas AMQP, capturas de Central y de Azure). "
        "**Límites:** una VM/región y N=51 por protocolo; AMQP no se probó en el ESP32 ni sobre WebSockets; el firmware ESP32 usa `setInsecure()` (simplificación de laboratorio); "
        "las latencias (~100 ms) están dominadas por el Hub, así que no se afirma ventaja de latencia de ningún protocolo.", tam=7.5)


# ------------------------------------------------------------------ PAGINA 2
doc.add_page_break()
titulo("5. Evidencia en IoT Central y en la VM (misma app lab4-medidordeclima-unab, plantilla Hobo MX-100 v3)")
cuadricula([
    ("12-central-amqp-datos-sin-procesar.jpg", "Fig. 2 — AMQP: datos sin procesar con el JSON expandido (temperature, humedad, illuminance); dispositivo Conectado.", (40, 68, 900, 480)),
    ("16-central-dashboard-3-protocolos.jpg", "Fig. 3 — Dashboard: las 3 variables de MQTT, AMQP y HTTPS en la misma app (gráficas y último valor).", (30, 55, 1375, 640)),
    ("13-central-https-datos-sin-procesar.jpg", "Fig. 4 — HTTPS: la telemetría llega cada 5 s, pero el dispositivo figura Desconectado (no hay sesión).", (40, 68, 900, 480)),
    ("18-central-explorador-datos.jpg", "Fig. 5 — Explorador de datos: Temperature por dispositivo (los 3 protocolos + ESP32).", (30, 40, 1510, 690)),
], 2, ANCHO)
cuadricula([
    ("02-amqp-tramas-proton.png", "Fig. 6 — Tramas AMQP reales (PN_TRACE_FRM): SASL, open, begin, attach $cbs, flow (crédito 50), put-token → accepted, attach de telemetría (completo: evidencias/02b)."),
    ("03-https-vm-log.png", "Fig. 7 — Log del cliente HTTPS en la VM: DPS por REST, POST y HTTP 204 por cada lectura (curl: evidencias/11)."),
    ("15-central-esp32-datos-sin-procesar.jpg", "Fig. 8 — ESP32 en Wokwi en Central: HTTPS (sin eventos de conexión) y MQTT (conectado/desconectado).", (40, 68, 900, 690)),
], 3, ANCHO, tam_pie=6.5)
parrafo("Las capturas de Central muestran hora de Bogotá (UTC−5); los logs de la VM, UTC (p. ej. 0:15 en Central = 05:15 en el log).", tam=6.5, color=GRIS, cursiva=True, despues=1)

titulo("6. Tabla comparativa MQTT · AMQP · HTTPS  (†: de la documentación, no medido aquí)")
tabla([
    ["Criterio", "MQTT 3.1.1 (8883)", "AMQP 1.0 (5671)", "HTTPS REST (443)"],
    ["Overhead por mensaje (payload 53 B)", "**365 B** (QoS 0: 221)", "438 B", "901 B (sin keep-alive 5 699)"],
    ["Latencia percibida (ACK p50)", "103 ms", "103 ms", "109 ms (199 sin keep-alive)"],
    ["Fiabilidad", "QoS 0/1 (QoS 2 se degrada a 1, Lab 3); paho reencola", "`accepted/rejected/released` por mensaje; reenvío por la app", "Solo lo que programe la app (cola, reintentos, timeouts)"],
    ["Renovar el SAS", "Reconectar", "**En caliente (CBS)** — medido", "Un SAS nuevo en cada petición"],
    ["Microcontrolador / Wokwi", "**Probado** ESP32 (PubSubClient; TLS 3,8 s)", "**No implementado**: sin cliente ligero estándar en Arduino-ESP32 (el SDK C de Azure sí, pesado†)", "**Probado** ESP32 (HTTPClient; 4,0 s + 162 ms)"],
    ["Facilidad de debug", "Media (`IOTC_DEBUG`: paquetes crudos)", "Baja–media (`PN_TRACE_FRM`: estructuras + payload binario)", "**Alta** (`curl -i`, códigos estándar)"],
    ["Paso de firewalls", "8883 a veces bloqueado → **MQTT-WS 443 probado** (370 B/msg)", "5671; AMQP-WS 443† (Proton Python no lo expone)", "**443**, atraviesa proxys"],
    ["Comandos / estado en Central", "Métodos directos: `setAlertLed` → 200 (medido) · *Conectado*", "Comandos: sí† · estado: *Conectado* (medido)", "Comandos: no† · estado: *Desconectado* (medido)"],
    ["Complejidad real (líneas de código)", "130", "**255** + compilar Proton con TLS en Linux", "117"],
    ["Caso de uso ideal", "Telemetría continua de dispositivos", "Plano de servicio, *pipelines* con contrapresión, gateways†", "Sensores que duermen minutos; integración y diagnóstico"],
], [3.6, 4.6, 5.3, 5.1], tam=7)

titulo("7. Recomendaciones para la plataforma propia (Labs 5–8)")
vinneta("**Dispositivo (Wokwi / ESP futuro): MQTT sobre TLS 8883, QoS 1.** Menos bytes (365 vs 438/901 B), presencia (*Conectado*), comandos y librería probada en el ESP32; "
        "renovar el SAS **reconectando antes de que venza**. Plan B tras firewalls estrictos: **MQTT-WS 443** (mismo modelo, mismos bytes por mensaje).", tam=8)
vinneta("**Entre servicios Azure y en la plataforma propia: AMQP 1.0** (Service Bus / Event Hubs): sin penalización de latencia (103 ms), con crédito para **contrapresión**, "
        "liquidación por mensaje y renovación de credenciales sin cortar el flujo; consumidores **idempotentes** (clave `message-id`).", tam=8)
vinneta("**Nos quedamos con dos protocolos: MQTT (borde) + AMQP (servicios).** Cada protocolo extra suma superficie de autenticación, de depuración y de código (AMQP ya costó 255 líneas y compilar Proton). "
        "HTTPS queda como **herramienta** (diagnóstico con `curl`, webhooks, dispositivos que duermen minutos), no como transporte permanente: 901 B/msg, sin presencia y sin comandos.", tam=8)

doc.save(SALIDA)
print("OK", SALIDA)
